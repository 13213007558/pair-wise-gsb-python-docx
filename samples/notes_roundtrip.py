"""Round-trip sample for footnote/endnote support in python-docx.

Builds a document that exercises the interesting cases:

- a footnote referenced twice (a "duplicate reference"),
- an image inside a footnote (relationship owned by the footnotes part),
- a hyperlink inside a footnote (relationship owned by the footnotes part),
- a table inside an endnote,

then saves, reloads, adds another footnote (proving ids continue without
reuse), and saves again. The output is `notes-roundtrip.docx` next to this
script. Run with:

    python samples/notes_roundtrip.py
"""

from __future__ import annotations

import os
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
IMAGE_PATH = os.path.join(HERE, "..", "tests", "test_files", "python-icon.png")
OUT_PATH = os.path.join(HERE, "notes-roundtrip.docx")


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """Add an external hyperlink to `paragraph`.

    The relationship is added to the part containing the paragraph, which is
    the footnotes part for a footnote paragraph, never the main document part.
    """
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_document() -> Document:
    document = Document()
    document.add_heading("Footnotes and endnotes round-trip sample", level=1)

    p1 = document.add_paragraph("The first claim is made here.")
    p2 = document.add_paragraph("The first claim is repeated here.")
    p3 = document.add_paragraph("An endnote-worthy remark.")

    # -- a footnote with rich content: text, an image, and a hyperlink --
    footnote = document.add_footnote(p1.runs[0], text="Evidence: ")
    footnote.paragraphs[0].add_run().add_picture(IMAGE_PATH, width=Inches(0.25))
    footnote.paragraphs[0].add_run(" (icon), source: ")
    add_hyperlink(footnote.paragraphs[0], "https://www.python.org", "python.org")

    # -- a second reference to the same footnote (duplicate reference) --
    footnote.add_reference(p2.runs[0])

    # -- an endnote containing a table --
    endnote = document.add_endnote(p3.runs[0], text="Summary of values:")
    table = endnote.add_table(rows=2, cols=2, width=Inches(3))
    table.cell(0, 0).text = "metric"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "answer"
    table.cell(1, 1).text = "42"

    return document


def main() -> None:
    document = build_document()
    document.save(OUT_PATH)

    # -- round trip: reload, add another footnote, save again --
    reloaded = Document(OUT_PATH)
    existing_ids = [f.footnote_id for f in reloaded.footnotes]
    p = reloaded.add_paragraph("A note added after reload.")
    added = reloaded.add_footnote(p.runs[0], text="Added on the second pass.")
    assert added.footnote_id not in existing_ids, "must not reuse an existing id"
    reloaded.save(OUT_PATH)

    # -- verify the final artifact --
    final = Document(OUT_PATH)
    footnote = final.footnotes.get(1)
    assert footnote is not None and len(footnote.references) == 2
    assert footnote.number == 1
    assert footnote.paragraphs[0].hyperlinks[0].address == "https://www.python.org"
    endnote = final.endnotes.get(1)
    assert endnote is not None and endnote.tables[0].cell(1, 1).text == "42"
    assert [f.footnote_id for f in final.footnotes] == [1, added.footnote_id]

    archive = zipfile.ZipFile(OUT_PATH)
    footnotes_rels = archive.read("word/_rels/footnotes.xml.rels").decode()
    document_rels = archive.read("word/_rels/document.xml.rels").decode()
    assert "relationships/image" in footnotes_rels
    assert "https://www.python.org" in footnotes_rels
    assert "relationships/image" not in document_rels
    assert "https://www.python.org" not in document_rels
    footnotes_xml = archive.read("word/footnotes.xml").decode()
    assert 'w:type="separator"' in footnotes_xml
    assert 'w:type="continuationSeparator"' in footnotes_xml

    print("OK: wrote", OUT_PATH)
    print("  footnotes:", [(f.footnote_id, f.number) for f in final.footnotes])
    print("  endnotes:", [(e.endnote_id, e.number) for e in final.endnotes])
    print("  footnote 1 references:", len(footnote.references))
    print("  footnote 1 text:", footnote.text)


if __name__ == "__main__":
    main()
