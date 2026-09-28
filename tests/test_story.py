# pyright: reportPrivateUsage=false

"""Test suite for the docx.story (unified story traversal) module."""

from __future__ import annotations

from collections import Counter

import pytest

import docx
from docx.enum.section import WD_HEADER_FOOTER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import nsmap
from docx.oxml.parser import parse_xml
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .unitutil.file import test_file

NSDECLS = " ".join('xmlns:%s="%s"' % item for item in sorted(nsmap.items()))


def _p(text: str) -> str:
    return '<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>' % text


def _tbl_xml(text: str) -> str:
    return (
        "<w:tbl>"
        '<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        '<w:tblGrid><w:gridCol w:w="1440"/></w:tblGrid>'
        "<w:tr><w:tc>"
        '<w:tcPr><w:tcW w:w="1440" w:type="dxa"/></w:tcPr>' + _p(text) + "</w:tc></w:tr></w:tbl>"
    )


_WPS_GRAPHIC = (
    "<a:graphic><a:graphicData"
    ' uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr/>'
    "<wps:txbx><w:txbxContent>%s</w:txbxContent></wps:txbx>"
    "<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>"
)


def _drawingml_anchor(body_xml: str, shape_id: int, name: str) -> str:
    """Floating DrawingML text box (`wp:anchor`) as a `w:drawing` element."""
    return (
        "<w:drawing>"
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"'
        ' relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="914400" cy="457200"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="%d" name="%s"/><wp:cNvGraphicFramePr/>'
        + _WPS_GRAPHIC
        + "</wp:anchor></w:drawing>"
    ) % (shape_id, name, body_xml)


def _drawingml_inline(body_xml: str, shape_id: int, name: str) -> str:
    """Inline DrawingML text box (`wp:inline`) as a `w:drawing` element."""
    return (
        "<w:drawing>"
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="914400" cy="457200"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:docPr id="%d" name="%s"/><wp:cNvGraphicFramePr/>'
        + _WPS_GRAPHIC
        + "</wp:inline></w:drawing>"
    ) % (shape_id, name, body_xml)


def _vml_pict(body_xml: str, name: str) -> str:
    """Legacy VML text box as a `w:pict` element."""
    return (
        '<w:pict><v:shape id="%s" type="#_x0000_t202"'
        ' style="width:72pt;height:36pt" stroked="f">'
        "<v:textbox><w:txbxContent>%s</w:txbxContent></v:textbox>"
        "</v:shape></w:pict>"
    ) % (name, body_xml)


def _run(inner_xml: str) -> str:
    return "<w:r %s>%s</w:r>" % (NSDECLS, inner_xml)


def _ac_run(choice_inner: str, fallback_inner: str) -> str:
    """`w:r` holding `mc:AlternateContent` with a DrawingML choice, VML fallback."""
    return (
        "<w:r %s><mc:AlternateContent>"
        '<mc:Choice Requires="wps">%s</mc:Choice>'
        "<mc:Fallback>%s</mc:Fallback></mc:AlternateContent></w:r>"
    ) % (NSDECLS, choice_inner, fallback_inner)


def _append_run(paragraph: Paragraph, run_xml: str) -> None:
    paragraph._p.append(parse_xml(run_xml))


@pytest.fixture(scope="module")
def populated_document(tmp_path_factory) -> docx.Document:
    """Document with nested tables, DrawingML/VML text boxes, shared headers."""
    document = docx.Document()

    # -- body flow content --
    document.add_paragraph("body-para-1")

    # -- floating DrawingML text box, duplicated as a legacy VML fallback --
    txbx_body = _p("txbx-para-1") + _tbl_xml("txbx-tbl-para")
    anchor_p = document.add_paragraph()
    _append_run(
        anchor_p,
        _ac_run(
            _drawingml_anchor(txbx_body, 1, "FloatingChoice"),
            _vml_pict(txbx_body, "FallbackShape"),
        ),
    )

    # -- AlternateContent whose branches carry distinguishable text --
    ac_p = document.add_paragraph()
    _append_run(
        ac_p,
        _ac_run(
            _drawingml_anchor(_p("ac-choice-para"), 2, "ChoiceOnly"),
            _vml_pict(_p("ac-fallback-para"), "FallbackOnly"),
        ),
    )

    # -- legacy VML-only text box (no AlternateContent wrapper) --
    vml_p = document.add_paragraph()
    _append_run(vml_p, _run(_vml_pict(_p("vml-txbx-para"), "VmlOnly")))

    # -- outer table with a nested table and a text box in another cell --
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("outer-cell-para")
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].add_run("inner-cell-para")
    cell_txbx_p = table.cell(1, 1).paragraphs[0]
    _append_run(cell_txbx_p, _run(_drawingml_anchor(_p("cell-txbx-para"), 3, "CellTxbx")))

    document.add_paragraph("body-para-2")

    # -- header and footer, each with its own part --
    section = document.sections[0]
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    header.paragraphs[0].add_run("header-para")
    footer = section.footer
    footer.paragraphs[0].add_run("footer-para")

    # -- inline DrawingML text box in the header holding a hyperlink and image --
    header_part = header.part
    link_rId = header_part.relate_to("https://example.com/", RT.HYPERLINK, is_external=True)
    image_rId, _image = header_part.get_or_add_image(test_file("python-icon.png"))
    header_anchor_p = header.add_paragraph()
    header_anchor_p._p.append(
        parse_xml(
            _run(
                _drawingml_inline(
                    '<w:p><w:hyperlink r:id="%s" w:history="1">'
                    "<w:r><w:t>header-link</w:t></w:r></w:hyperlink>"
                    '<w:r><w:drawing><wp:inline distT="0" distB="0"'
                    ' distL="0" distR="0">'
                    '<wp:extent cx="222250" cy="222250"/>'
                    '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
                    '<wp:docPr id="4" name="HeaderPicture"/><wp:cNvGraphicFramePr/>'
                    "<a:graphic><a:graphicData"
                    ' uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
                    "<pic:pic><pic:nvPicPr>"
                    '<pic:cNvPr id="4" name="python-icon.png"/><pic:cNvPicPr/>'
                    "</pic:nvPicPr><pic:blipFill>"
                    '<a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch>'
                    "</pic:blipFill><pic:spPr><a:xfrm>"
                    '<a:off x="0" y="0"/><a:ext cx="222250" cy="222250"/></a:xfrm>'
                    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
                    "</pic:pic></a:graphicData></a:graphic>"
                    "</wp:inline></w:drawing></w:r></w:p>" % (link_rId, image_rId),
                    5,
                    "HeaderTextBox",
                )
            )
        )
    )

    # -- second section explicitly referencing the same header/footer parts --
    document.add_section()
    first_sectPr = document.sections[0]._sectPr
    header_ref = first_sectPr.xpath("./w:headerReference")[0]
    footer_ref = first_sectPr.xpath("./w:footerReference")[0]
    second_sectPr = document.sections[1]._sectPr
    second_sectPr.add_headerReference(WD_HEADER_FOOTER.PRIMARY, header_ref.rId)
    second_sectPr.add_footerReference(WD_HEADER_FOOTER.PRIMARY, footer_ref.rId)

    path = tmp_path_factory.mktemp("story") / "populated.docx"
    document.save(path)
    return docx.Document(str(path))


def _item_by_text(items, text: str):
    matches = [item for item in items if getattr(item.content, "text", None) == text]
    assert len(matches) == 1, "expected exactly one %r item, got %d" % (text, len(matches))
    return matches[0]


class DescribeStory:
    """Unit-test suite for `docx.story` unified story traversal."""

    def it_visits_body_textboxes_and_nested_tables_once(self, populated_document):
        items = list(populated_document.stories.main.iter_items())
        paragraph_texts = [
            item.content.text for item in items if isinstance(item.content, Paragraph)
        ]
        counts = Counter(paragraph_texts)
        for text in (
            "body-para-1",
            "body-para-2",
            "txbx-para-1",
            "txbx-tbl-para",
            "vml-txbx-para",
            "outer-cell-para",
            "inner-cell-para",
            "cell-txbx-para",
            "ac-choice-para",
        ):
            assert counts[text] == 1, "%r appeared %d times" % (text, counts[text])
        # -- the rejected VML fallback content must not be extracted --
        assert "ac-fallback-para" not in counts
        # -- outer table, nested table, and text-box table, each once --
        assert sum(isinstance(item.content, Table) for item in items) == 3

    def it_traverses_in_xml_anchor_order(self, populated_document):
        paragraph_texts = [
            item.content.text
            for item in populated_document.stories.main.iter_items()
            if isinstance(item.content, Paragraph)
        ]
        expected = [
            "body-para-1",
            "",  # anchor paragraph of floating AC text box
            "txbx-para-1",
            "txbx-tbl-para",
            "",  # anchor paragraph of second AlternateContent
            "ac-choice-para",
            "",  # anchor paragraph of legacy VML text box
            "vml-txbx-para",
            "outer-cell-para",
            "inner-cell-para",
            "",  # trailing paragraph after nested table
            "",  # empty cell (0, 1)
            "",  # empty cell (1, 0)
            "",  # anchor paragraph of cell text box
            "cell-txbx-para",
            "body-para-2",
            "",  # section-break paragraph added by `add_section()`
        ]
        assert paragraph_texts == expected

    def it_records_the_nested_container_path(self, populated_document):
        items = list(populated_document.stories.main.iter_items())

        inner_cell_item = _item_by_text(items, "inner-cell-para")
        assert tuple(c.kind for c in inner_cell_item.path) == (
            "body",
            "cell",
            "cell",
        )
        assert inner_cell_item.path[0].proxy is populated_document._body
        assert isinstance(inner_cell_item.path[1].proxy, _Cell)
        assert inner_cell_item.path[1].index == (0, 0)
        assert inner_cell_item.path[2].index == (0, 0)
        assert inner_cell_item.anchor is None

        txbx_cell_item = _item_by_text(items, "txbx-tbl-para")
        assert tuple(c.kind for c in txbx_cell_item.path) == (
            "body",
            "textbox",
            "cell",
        )

        cell_txbx_item = _item_by_text(items, "cell-txbx-para")
        assert tuple(c.kind for c in cell_txbx_item.path) == (
            "body",
            "cell",
            "textbox",
        )
        assert cell_txbx_item.path[1].index == (1, 1)

    def it_associates_textboxes_with_their_xml_anchor(self, populated_document):
        items = list(populated_document.stories.main.iter_items())

        body_item = _item_by_text(items, "body-para-1")
        choice_item = _item_by_text(items, "txbx-para-1")
        vml_item = _item_by_text(items, "vml-txbx-para")
        ac_item = _item_by_text(items, "ac-choice-para")
        cell_item = _item_by_text(items, "cell-txbx-para")

        assert body_item.anchor is None
        assert choice_item.anchor.is_floating is True
        assert vml_item.anchor.is_floating is True
        assert ac_item.anchor.is_floating is True
        assert cell_item.anchor.is_floating is True
        # -- XML anchor order, not visual layout order --
        assert [
            choice_item.anchor.order,
            ac_item.anchor.order,
            vml_item.anchor.order,
            cell_item.anchor.order,
        ] == [0, 1, 2, 3]
        assert choice_item.anchor.paragraph.text == ""
        assert _item_by_text(items, "body-para-2").anchor is None

    def it_selects_explicitly_supported_ac_branches(self, populated_document):
        default_items = list(populated_document.stories.main.iter_items())
        default_texts = [
            item.content.text for item in default_items if isinstance(item.content, Paragraph)
        ]
        assert "ac-choice-para" in default_texts
        assert "ac-fallback-para" not in default_texts

        fallback_items = list(populated_document.stories.main.iter_items(capabilities=frozenset()))
        fallback_texts = [
            item.content.text for item in fallback_items if isinstance(item.content, Paragraph)
        ]
        assert "ac-fallback-para" in fallback_texts
        assert "ac-choice-para" not in fallback_texts
        # -- content is still extracted exactly once via the fallback branch --
        assert fallback_texts.count("txbx-para-1") == 1
        assert fallback_texts.count("txbx-tbl-para") == 1

    def it_returns_header_and_footer_as_separate_stories(self, populated_document):
        stories = populated_document.stories
        kinds = [story.kind for story in stories]

        assert kinds == ["body", "header", "footer"]
        header_story = stories[1]
        footer_story = stories[2]

        # -- the same part referenced by two sections yields a single story --
        assert header_story.sections == (0, 1)
        assert footer_story.sections == (0, 1)
        assert [ref.header_footer_type for ref in header_story.references] == [
            WD_HEADER_FOOTER.PRIMARY,
            WD_HEADER_FOOTER.PRIMARY,
        ]
        assert len(stories) == 3

        header_part = populated_document.sections[0].header.part
        footer_part = populated_document.sections[1].footer.part
        assert stories.for_part(populated_document.part) is stories.main
        assert stories.for_part(header_part) is header_story
        assert stories.for_part(footer_part) is footer_story

    def it_carries_the_owning_part_for_header_textbox_content(self, populated_document):
        header_part = populated_document.sections[0].header.part
        header_story = populated_document.stories.for_part(header_part)
        items = list(header_story.iter_items())

        header_item = _item_by_text(items, "header-para")
        link_item = _item_by_text(items, "header-link")

        assert header_item.part is header_part
        assert link_item.part is header_part
        assert tuple(c.kind for c in link_item.path) == ("header", "textbox")
        assert link_item.anchor.is_floating is False

        # -- hyperlink resolves against the header part, not the document part --
        hyperlink = link_item.content.hyperlinks[0]
        assert hyperlink.address == "https://example.com/"
        image_rIds = [
            rId
            for rId, relationship in header_part.rels.items()
            if relationship.reltype == RT.IMAGE
        ]
        assert image_rIds, "header text box image relationship must live on header part"
        document_image_rIds = [
            rId
            for rId, relationship in populated_document.part.rels.items()
            if relationship.reltype == RT.IMAGE
        ]
        assert document_image_rIds == []

        all_paragraph_texts = [
            item.content.text
            for item in populated_document.stories.iter_items()
            if isinstance(item.content, Paragraph)
        ]
        assert all_paragraph_texts.count("header-para") == 1
        assert all_paragraph_texts.count("footer-para") == 1
        assert all_paragraph_texts.count("txbx-para-1") == 1
