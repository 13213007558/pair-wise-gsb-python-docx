"""Test suite for list-numbering resolution (docx.numbering).

The fixture document mixes, in a single document:
- a multilevel list that continues numbering across interruptions (same template),
- a second numbering instance on the same abstractNum that explicitly restarts
  via w:startOverride,
- a paragraph that jumps directly into a third-level list,
- list paragraphs inside a table, participating in document order,
- style-linked numbering (both style -> numPr and lvl -> pStyle links),
- w:lvlRestart behavior (default restart and "never restart"),
- a Chinese-counting format ("第十二条") and an unsupported format (w:ordinal).
"""

from __future__ import annotations

from lxml import etree

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.numbering import SUPPORTED_NUM_FMTS
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

NUMBERING_XML = (
    "<w:numbering {ns}>"
    # -- multilevel decimal template shared by numId 1, 2, and 7 --
    '<w:abstractNum w:abstractNumId="10">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:lvlText w:val="%1."/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:lvlText w:val="%1.%2)"/></w:lvl>'
    '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:lvlText w:val="(%1.%2.%3)"/></w:lvl>'
    "</w:abstractNum>"
    # -- Chinese-counting clauses starting at 12 --
    '<w:abstractNum w:abstractNumId="11">'
    '<w:lvl w:ilvl="0"><w:start w:val="12"/><w:numFmt w:val="chineseCountingThousand"/>'
    '<w:lvlText w:val="第%1条"/></w:lvl>'
    "</w:abstractNum>"
    # -- lowerLetter sub-level that never restarts (lvlRestart val=0) --
    '<w:abstractNum w:abstractNumId="12">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:lvlText w:val="%1."/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
    '<w:lvlRestart w:val="0"/><w:lvlText w:val="(%2)"/></w:lvl>'
    "</w:abstractNum>"
    # -- upperRoman/lowerLetter with default restart behavior --
    '<w:abstractNum w:abstractNumId="13">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperRoman"/>'
    '<w:lvlText w:val="%1."/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
    '<w:lvlText w:val="%2)"/></w:lvl>'
    "</w:abstractNum>"
    # -- unsupported number format --
    '<w:abstractNum w:abstractNumId="14">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="ordinal"/>'
    '<w:lvlText w:val="%1."/></w:lvl>'
    "</w:abstractNum>"
    # -- level linked to the "ClauseHead" paragraph style --
    '<w:abstractNum w:abstractNumId="15">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
    '<w:pStyle w:val="ClauseHead"/><w:lvlText w:val="%1."/></w:lvl>'
    "</w:abstractNum>"
    # -- bullet list --
    '<w:abstractNum w:abstractNumId="16">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
    '<w:lvlText w:val="•"/></w:lvl>'
    "</w:abstractNum>"
    # -- concrete numbering instances --
    '<w:num w:numId="1"><w:abstractNumId w:val="10"/></w:num>'
    '<w:num w:numId="2"><w:abstractNumId w:val="10"/>'
    '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride></w:num>'
    '<w:num w:numId="3"><w:abstractNumId w:val="11"/></w:num>'
    '<w:num w:numId="4"><w:abstractNumId w:val="12"/></w:num>'
    '<w:num w:numId="5"><w:abstractNumId w:val="13"/></w:num>'
    '<w:num w:numId="6"><w:abstractNumId w:val="14"/></w:num>'
    '<w:num w:numId="7"><w:abstractNumId w:val="10"/></w:num>'
    '<w:num w:numId="8"><w:abstractNumId w:val="15"/></w:num>'
    '<w:num w:numId="9"><w:abstractNumId w:val="16"/></w:num>'
    "</w:numbering>"
)


def _replace_numbering(doc, numbering_xml):
    numbering_elm = doc.part.numbering_part.element
    for child in list(numbering_elm):
        numbering_elm.remove(child)
    for child in parse_xml(numbering_xml.replace("{ns}", nsdecls("w"))):
        numbering_elm.append(child)


def _add_num_pr(paragraph, num_id, ilvl):
    numPr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = ilvl
    numPr.get_or_add_numId().val = num_id


def _add_list_paragraph(container, text, num_id, ilvl, style=None):
    paragraph = container.add_paragraph(text, style=style)
    if num_id is not None:
        _add_num_pr(paragraph, num_id, ilvl)
    return paragraph


def _mixed_list_document():
    """Document mixing continued numbering, explicit restart, level-jumping,
    table paragraphs, style links, lvlRestart, and assorted number formats."""
    doc = Document()
    _replace_numbering(doc, NUMBERING_XML)

    doc.styles.add_style("ClauseHead", WD_STYLE_TYPE.PARAGRAPH)
    list_para_style = doc.styles.add_style("ListPara", WD_STYLE_TYPE.PARAGRAPH)
    numPr = list_para_style.element.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 1
    numPr.get_or_add_numId().val = 1

    # -- multilevel list that continues across interruptions --
    _add_list_paragraph(doc, "item one", 1, 0)
    _add_list_paragraph(doc, "item one point one", 1, 1)
    doc.add_paragraph("plain paragraph, not a list item")
    # -- list paragraphs inside a table participate in document order --
    table = doc.add_table(rows=1, cols=2)
    _add_list_paragraph(table.cell(0, 0), "in table one", 1, 0)
    _add_list_paragraph(table.cell(0, 1), "in table two", 1, 0)
    _add_list_paragraph(doc, "item four", 1, 0)
    # -- explicit restart on the same template via startOverride --
    _add_list_paragraph(doc, "restarted five", 2, 0)
    _add_list_paragraph(doc, "restarted six", 2, 0)
    # -- jump directly into a third-level list --
    _add_list_paragraph(doc, "deep item", 7, 2)
    # -- Chinese-counting clauses --
    _add_list_paragraph(doc, "clause twelve", 3, 0)
    _add_list_paragraph(doc, "clause thirteen", 3, 0)
    # -- sub-level that never restarts (lvlRestart val=0) --
    _add_list_paragraph(doc, "first", 4, 0)
    _add_list_paragraph(doc, "first a", 4, 1)
    _add_list_paragraph(doc, "second", 4, 0)
    _add_list_paragraph(doc, "second b", 4, 1)
    # -- sub-level with default restart behavior --
    _add_list_paragraph(doc, "roman one", 5, 0)
    _add_list_paragraph(doc, "roman one a", 5, 1)
    _add_list_paragraph(doc, "roman two", 5, 0)
    _add_list_paragraph(doc, "roman two a", 5, 1)
    # -- unsupported number format --
    _add_list_paragraph(doc, "ordinal one", 6, 0)
    _add_list_paragraph(doc, "ordinal two", 6, 0)
    # -- numbering linked from w:lvl/w:pStyle --
    _add_list_paragraph(doc, "clause head one", None, None, style="ClauseHead")
    _add_list_paragraph(doc, "clause head two", None, None, style="ClauseHead")
    # -- numbering linked from style w:pPr/w:numPr --
    _add_list_paragraph(doc, "styled sub item", None, None, style="ListPara")
    # -- bullet --
    _add_list_paragraph(doc, "bullet item", 9, 0)
    # -- numId 0 and unknown numId are not list items --
    _add_list_paragraph(doc, "numbering removed", 0, 0)
    _add_list_paragraph(doc, "unknown numbering", 99, 0)
    return doc


class DescribeIterListItems:
    def it_yields_each_list_paragraph_in_document_order(self):
        items = list(_mixed_list_document().iter_list_items())

        assert [item.label for item in items] == [
            "1.",           # numId 1, continuing multilevel list
            "1.1)",
            "2.",           # table cell (0, 0), in document order
            "3.",           # table cell (0, 1)
            "4.",           # continues after the table
            "5.",           # numId 2, explicit restart via startOverride
            "6.",
            "(1.1.1)",      # numId 7, jumping directly into the third level
            "第十二条",
            "第十三条",
            "1.",           # numId 4
            "(a)",
            "2.",
            "(b)",          # lvlRestart val=0: sub-level does not restart
            "I.",           # numId 5
            "a)",
            "II.",
            "a)",           # default: sub-level restarts after level 0
            None,           # numId 6, unsupported "ordinal" format
            None,
            "1.",           # numId 8, linked via w:lvl/w:pStyle
            "2.",
            "4.1)",         # numId 1 again, linked via style w:numPr
            "\u2022",       # bullet
        ]

    def it_reports_level_value_and_num_id_for_each_item(self):
        items = list(_mixed_list_document().iter_list_items())

        assert [item.num_id for item in items] == [
            1, 1, 1, 1, 1, 2, 2, 7, 3, 3, 4, 4, 4, 4, 5, 5, 5, 5, 6, 6, 8, 8, 1, 9,
        ]
        assert [item.level for item in items] == [
            0, 1, 0, 0, 0, 0, 0, 2, 0, 0, 0, 1, 0, 1, 0, 1, 0, 1, 0, 0, 0, 0, 1, 0,
        ]
        assert [item.value for item in items] == [
            1, 1, 2, 3, 4, 5, 6, 1, 12, 13, 1, 1, 2, 2, 1, 1, 2, 1, 1, 2, 1, 2, 1, None,
        ]

    def it_tracks_counters_independently_for_each_num_id(self):
        items = list(_mixed_list_document().iter_list_items())
        by_num_id = {}
        for item in items:
            by_num_id.setdefault(item.num_id, []).append(item.value)

        # -- numId 2 restarts explicitly at 5 rather than continuing numId 1,
        #    and numId 7 starts its own counters on the same template --
        assert by_num_id[1] == [1, 1, 2, 3, 4, 1]
        assert by_num_id[2] == [5, 6]
        assert by_num_id[7] == [1]

    def it_initializes_ancestor_counters_when_jumping_to_a_deeper_level(self):
        items = list(_mixed_list_document().iter_list_items())
        deep_item = next(item for item in items if item.num_id == 7)

        assert deep_item.level == 2
        assert deep_item.value == 1
        assert deep_item.label == "(1.1.1)"

    def it_expands_upper_level_references_in_lvl_text(self):
        items = list(_mixed_list_document().iter_list_items())
        styled_sub_item = items[-2]

        # -- "%1.%2)" expands the current level-0 counter of numId 1 --
        assert styled_sub_item.num_id == 1
        assert styled_sub_item.level == 1
        assert styled_sub_item.label == "4.1)"

    def it_reports_unsupported_formats_without_coercing_to_decimal(self):
        items = list(_mixed_list_document().iter_list_items())
        ordinal_items = [item for item in items if item.num_id == 6]

        assert "ordinal" not in SUPPORTED_NUM_FMTS
        assert [item.is_num_fmt_supported for item in ordinal_items] == [False, False]
        assert [item.num_fmt for item in ordinal_items] == ["ordinal", "ordinal"]
        assert [item.label for item in ordinal_items] == [None, None]
        # -- counters still advance so position information is not lost --
        assert [item.value for item in ordinal_items] == [1, 2]

    def it_marks_bullets_with_a_literal_label_and_no_value(self):
        items = list(_mixed_list_document().iter_list_items())
        bullet_item = items[-1]

        assert bullet_item.num_fmt == "bullet"
        assert bullet_item.label == "\u2022"
        assert bullet_item.value is None
        assert bullet_item.is_num_fmt_supported is True

    def it_links_each_item_to_its_source_paragraph(self):
        doc = _mixed_list_document()
        items = list(doc.iter_list_items())

        assert items[0].paragraph.text == "item one"
        assert items[2].paragraph.text == "in table one"
        assert items[0].paragraph._p is doc.paragraphs[0]._p

    def it_does_not_modify_document_or_numbering_xml(self):
        doc = _mixed_list_document()
        numbering_elm = doc.part.numbering_part.element
        before = (
            etree.tostring(doc.element),
            etree.tostring(numbering_elm),
            [p.text for p in doc.paragraphs],
        )

        list(doc.iter_list_items())

        assert etree.tostring(doc.element) == before[0]
        assert etree.tostring(numbering_elm) == before[1]
        # -- automatic numbering is not converted to literal paragraph text --
        assert [p.text for p in doc.paragraphs] == before[2]

    def it_yields_nothing_when_there_is_no_numbering_part(self):
        doc = Document()
        rels = doc.part.rels
        for rId in [rId for rId in rels if "numbering" in rels[rId].reltype]:
            rels.pop(rId)

        assert list(doc.iter_list_items()) == []
