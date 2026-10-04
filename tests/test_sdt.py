# pyright: reportPrivateUsage=false

"""Unit test suite for the 'docx.sdt' module."""

from __future__ import annotations

import io

import pytest

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import parse_xml
from docx.parts.customxml import CustomXmlPart
from docx.sdt import (
    BindingErrorKind,
    ContentControlRefreshError,
    ControlType,
    SimpleXPath,
    XPathNotSupportedError,
    parse_prefix_mappings,
)

STORE_ITEM_ID = "{11111111-2222-3333-4444-555555555555}"
NS = "http://example.com/ns"
PFX = "xmlns:ns='%s'" % NS
DS_NS = "http://schemas.openxmlformats.org/officeDocument/2006/customXml"


def store_xml(*item_names: str) -> bytes:
    items = "".join("<ns:item><ns:name>%s</ns:name></ns:item>" % n for n in item_names)
    return (
        '<ns:root xmlns:ns="%s">'
        "<ns:name>Alice</ns:name>"
        "<ns:approved>1</ns:approved>"
        "<ns:code ns:lang='en'>X1</ns:code>"
        "%s"
        "</ns:root>" % (NS, items)
    ).encode("utf-8")


def add_store(document, xml_data: bytes = store_xml(), item_id: str = STORE_ITEM_ID):
    part = document.part
    cx_part = CustomXmlPart(
        PackURI("/customXml/item99.xml"), "application/xml", parse_xml(xml_data), part.package
    )
    part.relate_to(cx_part, RT.CUSTOM_XML)
    props_xml = (
        "<ds:datastoreItem xmlns:ds='%s' ds:itemID='%s'/>" % (DS_NS, item_id)
    ).encode("utf-8")
    props_part = Part(
        PackURI("/customXml/itemProps99.xml"),
        CT.OFC_CUSTOM_XML_PROPERTIES,
        props_xml,
        part.package,
    )
    cx_part.relate_to(props_part, RT.CUSTOM_XML_PROPS)
    return cx_part


def add_to_body(document, *fragments: str) -> None:
    body = document.part.element.body
    sectPr = body.find(qn("w:sectPr"))
    for fragment in fragments:
        elm = parse_xml(fragment)
        if sectPr is not None:
            sectPr.addprevious(elm)
        else:
            body.append(elm)


def text_sdt(sdt_id: int, alias: str, xpath: str | None, text: str = "OLD") -> str:
    binding = (
        '<w:dataBinding w:storeItemID="%s" w:xpath="%s" w:prefixMappings="%s"/>'
        % (STORE_ITEM_ID, xpath, PFX)
        if xpath
        else ""
    )
    return (
        "<w:p %s><w:sdt><w:sdtPr>"
        '<w:alias w:val="%s"/><w:id w:val="%d"/>%s<w:text/>'
        "</w:sdtPr><w:sdtContent><w:r><w:t>%s</w:t></w:r></w:sdtContent>"
        "</w:sdt></w:p>" % (nsdecls("w", "w14", "w15"), alias, sdt_id, binding, text)
    )


def checkbox_sdt(sdt_id: int, alias: str, xpath: str) -> str:
    return (
        "<w:p %s><w:sdt><w:sdtPr>"
        '<w:alias w:val="%s"/><w:id w:val="%d"/>'
        '<w:dataBinding w:storeItemID="%s" w:xpath="%s" w:prefixMappings="%s"/>'
        "<w14:checkbox>"
        '<w14:checked w14:val="0"/>'
        '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/>'
        "</w14:checkbox>"
        "</w:sdtPr><w:sdtContent><w:r><w:t>☐</w:t></w:r></w:sdtContent>"
        "</w:sdt></w:p>"
        % (nsdecls("w", "w14", "w15"), alias, sdt_id, STORE_ITEM_ID, xpath, PFX)
    )


def section_sdt(sdt_id: int, child_id: int, child_xpath: str) -> str:
    return (
        "<w:sdt %s><w:sdtPr>"
        '<w:alias w:val="Items"/><w:id w:val="%d"/>'
        '<w:dataBinding w:storeItemID="%s" w:xpath="/ns:root/ns:item" w:prefixMappings="%s"/>'
        "<w15:repeatingSection/>"
        "</w:sdtPr><w:sdtContent>"
        "<w:p><w:sdt><w:sdtPr>"
        '<w:alias w:val="Item name"/><w:id w:val="%d"/>'
        '<w:dataBinding w:storeItemID="%s" w:xpath="%s" w:prefixMappings="%s"/>'
        "<w:text/>"
        "</w:sdtPr><w:sdtContent><w:r><w:t>?</w:t></w:r></w:sdtContent>"
        "</w:sdt></w:p>"
        "</w:sdtContent></w:sdt>"
        % (
            nsdecls("w", "w14", "w15"),
            sdt_id,
            STORE_ITEM_ID,
            PFX,
            child_id,
            STORE_ITEM_ID,
            child_xpath,
            PFX,
        )
    )


def body_sdts(document):
    return [
        e
        for e in document.part.element.body.iter(qn("w:sdt"))
        if e.getparent() is document.part.element.body
    ]


class DescribeParsePrefixMappings:
    def it_parses_single_and_double_quoted_mappings(self):
        mappings = parse_prefix_mappings('xmlns:a=\'uri-a\' xmlns:b="uri-b"')
        assert mappings == {"a": "uri-a", "b": "uri-b"}

    def it_returns_an_empty_mapping_for_empty_input(self):
        assert parse_prefix_mappings("") == {}
        assert parse_prefix_mappings(None) == {}


class DescribeSimpleXPath:
    def it_selects_elements_by_absolute_path(self):
        root = parse_xml(store_xml())
        path = SimpleXPath.parse("/ns:root/ns:name", {"ns": NS})
        assert [e.text for e in path.select(root)] == ["Alice"]

    def it_honors_one_based_positional_predicates(self):
        root = parse_xml(store_xml("a", "b", "c"))
        path = SimpleXPath.parse("/ns:root/ns:item[2]/ns:name", {"ns": NS})
        assert [e.text for e in path.select(root)] == ["b"]

    def it_selects_attribute_values(self):
        root = parse_xml(store_xml())
        path = SimpleXPath.parse("/ns:root/ns:code/@ns:lang", {"ns": NS})
        assert path.select(root) == ["en"]

    def it_evaluates_relative_paths_from_the_context_node(self):
        root = parse_xml(store_xml("a", "b"))
        item = root.findall("{%s}item" % NS)[1]
        path = SimpleXPath.parse("ns:name", {"ns": NS})
        assert [e.text for e in path.select(item)] == ["b"]

    def it_rebases_an_absolute_path_onto_a_section_path(self):
        section = SimpleXPath.parse("/ns:root/ns:item", {"ns": NS})
        child = SimpleXPath.parse("/ns:root/ns:item[3]/ns:name", {"ns": NS})
        rebased = child.relative_to(section)
        assert rebased is not None
        root = parse_xml(store_xml("a", "b", "c"))
        item = root.findall("{%s}item" % NS)[1]
        assert [e.text for e in rebased.select(item)] == ["b"]

    def it_returns_none_when_a_path_does_not_extend_the_section_path(self):
        section = SimpleXPath.parse("/ns:root/ns:item", {"ns": NS})
        other = SimpleXPath.parse("/ns:root/ns:name", {"ns": NS})
        assert other.relative_to(section) is None

    @pytest.mark.parametrize(
        "xpath",
        ["//ns:name", "/ns:root//ns:name", "/ns:root/*", "/ns:root/ns:item[last()]",
         "/ns:root/ns:item[@a='b']", "..", "/ns:root/ns:item[0]", ""],
    )
    def it_rejects_xpaths_outside_the_supported_subset(self, xpath: str):
        with pytest.raises(XPathNotSupportedError):
            SimpleXPath.parse(xpath, {"ns": NS})

    def it_rejects_unknown_namespace_prefixes(self):
        with pytest.raises(XPathNotSupportedError):
            SimpleXPath.parse("/xx:root", {"ns": NS})


class DescribeCustomXmlStores:
    def it_discovers_stores_by_their_item_id(self):
        document = Document()
        add_store(document)
        stores = document.custom_xml_stores
        assert STORE_ITEM_ID in stores
        assert stores[STORE_ITEM_ID].element.tag == "{%s}root" % NS

    def it_gets_and_sets_values_by_xpath(self):
        document = Document()
        add_store(document)
        store = document.custom_xml_stores[STORE_ITEM_ID]
        store.set_value("/ns:root/ns:name", "Bob", PFX)
        assert store.get_value("/ns:root/ns:name", PFX) == "Bob"
        store.set_value("/ns:root/ns:code/@ns:lang", "fr", PFX)
        assert store.get_value("/ns:root/ns:code/@ns:lang", PFX) == "fr"

    def it_raises_when_set_value_matches_an_unexpected_node_count(self):
        document = Document()
        add_store(document)
        store = document.custom_xml_stores[STORE_ITEM_ID]
        with pytest.raises(ValueError):
            store.set_value("/ns:root/ns:nope", "x", PFX)


class DescribeContentControls:
    def it_lists_the_controls_of_the_document(self):
        document = Document()
        add_store(document)
        add_to_body(document, text_sdt(10, "Name", "/ns:root/ns:name"))
        controls = list(document.content_controls)
        assert len(controls) == 1
        assert controls[0].alias == "Name"
        assert controls[0].id == 10
        assert controls[0].control_type == ControlType.PLAIN_TEXT
        assert controls[0].text == "OLD"


class DescribePlainTextRefresh:
    def it_updates_bound_controls_and_leaves_unbound_ones_alone(self):
        document = Document()
        add_store(document)
        add_to_body(
            document,
            text_sdt(10, "Name", "/ns:root/ns:name"),
            text_sdt(11, "Static", None, "KEEP"),
        )
        document.content_controls.refresh()
        controls = list(document.content_controls)
        assert controls[0].text == "Alice"
        assert controls[1].text == "KEEP"

    def it_updates_all_controls_bound_to_the_same_data(self):
        document = Document()
        add_store(document)
        add_to_body(
            document,
            text_sdt(10, "Name 1", "/ns:root/ns:name"),
            text_sdt(11, "Name 2", "/ns:root/ns:name"),
        )
        document.content_controls.refresh()
        assert [c.text for c in document.content_controls] == ["Alice", "Alice"]

    def it_separates_planning_from_applying(self):
        document = Document()
        add_store(document)
        add_to_body(document, text_sdt(10, "Name", "/ns:root/ns:name"))
        plan = document.content_controls.plan_refresh()
        assert plan.ok
        assert len(plan.actions) == 1
        assert plan.actions[0].kind == "set-text"
        assert list(document.content_controls)[0].text == "OLD"
        plan.apply()
        assert list(document.content_controls)[0].text == "Alice"


class DescribeCheckboxRefresh:
    def it_checks_and_unchecks_from_boolean_like_data(self):
        document = Document()
        add_store(document)
        add_to_body(document, checkbox_sdt(20, "Approved", "/ns:root/ns:approved"))
        document.content_controls.refresh()
        sdt = list(document.content_controls)[0]
        assert sdt.control_type == ControlType.CHECKBOX
        assert sdt.text == "☒"
        checked = document.part.element.body.find(
            ".//" + qn("w14:checkbox") + "/" + qn("w14:checked")
        )
        assert checked.get(qn("w14:val")) == "1"

    def it_reports_a_type_mismatch_for_non_boolean_data(self):
        document = Document()
        add_store(document, store_xml().replace(b">1<", b">maybe<"))
        add_to_body(document, checkbox_sdt(20, "Approved", "/ns:root/ns:approved"))
        plan = document.content_controls.plan_refresh()
        assert not plan.ok
        assert plan.errors[0].kind == BindingErrorKind.TYPE_MISMATCH
        assert "Approved" in plan.errors[0].control
        with pytest.raises(ContentControlRefreshError):
            plan.apply()
        assert list(document.content_controls)[0].text == "☐"


class DescribeRefreshErrors:
    def it_reports_unsupported_xpaths_before_committing(self):
        document = Document()
        add_store(document)
        add_to_body(document, text_sdt(10, "Name", "//ns:name"))
        plan = document.content_controls.plan_refresh()
        assert plan.errors[0].kind == BindingErrorKind.UNSUPPORTED_XPATH
        assert "Name" in plan.errors[0].control
        with pytest.raises(ContentControlRefreshError):
            plan.apply()

    def it_reports_an_unexpected_result_count(self):
        document = Document()
        add_store(document, store_xml("a", "b"))
        add_to_body(
            document,
            text_sdt(10, "Missing", "/ns:root/ns:nope"),
            text_sdt(11, "Many", "/ns:root/ns:item"),
        )
        plan = document.content_controls.plan_refresh()
        kinds = [e.kind for e in plan.errors]
        assert kinds == [BindingErrorKind.RESULT_COUNT, BindingErrorKind.RESULT_COUNT]
        assert "Missing" in plan.errors[0].control
        assert "Many" in plan.errors[1].control

    def it_reports_a_missing_data_store(self):
        document = Document()
        add_store(document)
        fragment = text_sdt(10, "Name", "/ns:root/ns:name").replace(
            STORE_ITEM_ID, "{NO-SUCH-STORE}"
        )
        add_to_body(document, fragment)
        plan = document.content_controls.plan_refresh()
        assert plan.errors[0].kind == BindingErrorKind.STORE_NOT_FOUND
        assert "Name" in plan.errors[0].control


class DescribeRepeatingSectionRefresh:
    def it_expands_the_section_to_one_control_per_data_node(self):
        document = Document()
        add_store(document, store_xml("a", "b", "c"))
        add_to_body(document, section_sdt(30, 31, "ns:name"))
        plan = document.content_controls.plan_refresh()
        assert plan.ok
        assert [a.kind for a in plan.actions] == ["expand-section"]
        plan.apply()
        sections = body_sdts(document)
        assert len(sections) == 3
        texts = [s.display_text() for s in sections]
        assert texts == ["a", "b", "c"]

    def it_assigns_unique_ids_to_generated_copies(self):
        document = Document()
        add_store(document, store_xml("a", "b", "c"))
        add_to_body(document, section_sdt(30, 31, "ns:name"))
        document.content_controls.refresh()
        ids = document.part.element.xpath(".//w:sdtPr/w:id/@w:val")
        assert len(ids) == len(set(ids))
        items = body_sdts(document)[1:]
        assert all(item.is_repeating_section_item for item in items)

    def it_rebases_word_style_absolute_child_xpaths_onto_each_item(self):
        document = Document()
        add_store(document, store_xml("a", "b"))
        add_to_body(document, section_sdt(30, 31, "/ns:root/ns:item[1]/ns:name"))
        document.content_controls.refresh()
        assert [s.display_text() for s in body_sdts(document)] == ["a", "b"]

    def it_does_not_accumulate_copies_when_refreshed_twice(self):
        document = Document()
        add_store(document, store_xml("a", "b", "c"))
        add_to_body(document, section_sdt(30, 31, "ns:name"))
        document.content_controls.refresh()
        document.content_controls.refresh()
        assert len(body_sdts(document)) == 3
        assert [s.display_text() for s in body_sdts(document)] == ["a", "b", "c"]

    def it_shrinks_back_when_the_data_has_fewer_nodes(self):
        document = Document()
        cx_part = add_store(document, store_xml("a", "b", "c"))
        add_to_body(document, section_sdt(30, 31, "ns:name"))
        document.content_controls.refresh()
        root = cx_part.element
        root.remove(root.findall("{%s}item" % NS)[2])
        document.content_controls.refresh()
        assert [s.display_text() for s in body_sdts(document)] == ["a", "b"]

    def it_reports_zero_matching_nodes_as_a_result_count_error(self):
        document = Document()
        add_store(document)
        add_to_body(document, section_sdt(30, 31, "ns:name"))
        plan = document.content_controls.plan_refresh()
        assert plan.errors[0].kind == BindingErrorKind.RESULT_COUNT
        assert "Items" in plan.errors[0].control


class DescribeRoundTrip:
    def it_preserves_stores_and_controls_across_save_and_open(self):
        document = Document()
        add_store(document, store_xml("a", "b"))
        add_to_body(
            document,
            text_sdt(10, "Name", "/ns:root/ns:name"),
            section_sdt(30, 31, "ns:name"),
        )
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        reopened = Document(stream)
        store = reopened.custom_xml_stores[STORE_ITEM_ID]
        store.set_value("/ns:root/ns:name", "Carol", PFX)
        reopened.content_controls.refresh()
        controls = list(reopened.content_controls)
        assert controls[0].text == "Carol"
        assert [s.display_text() for s in body_sdts(reopened)] == ["a", "b"]
