"""Integration tests for tracked-revision enumeration and processing."""

from __future__ import annotations

from io import BytesIO

import pytest

import docx
from docx.enum.revision import WD_REVISION_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.exceptions import UnsupportedRevisionError
from docx.oxml import parse_xml
from docx.oxml.revisions import apply_revision_actions


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _document(xml: str):
    document = docx.Document()
    element = parse_xml(f'<w:document xmlns:w="{W}">{xml}</w:document>')
    document._part._element = element
    document._element = element
    return document


def _reopen(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return docx.Document(stream)


class DescribeRevisions:
    def it_enumerates_supported_revision_types_with_metadata(self):
        document = _document(
            "<w:body>"
            '<w:p>'
            '<w:pPr><w:jc w:val="right"/>'
            '<w:pPrChange w:id="5" w:author="Carol">'
            '<w:pPr><w:jc w:val="left"/></w:pPr></w:pPrChange>'
            "</w:pPr>"
            '<w:ins w:id="1" w:author="Alice"><w:r><w:t>A</w:t></w:r></w:ins>'
            '<w:del w:id="2" w:author="Bob"><w:r><w:delText>B</w:delText></w:r></w:del>'
            '<w:r><w:rPr><w:b/>'
            '<w:rPrChange w:id="6" w:author="Carol"><w:rPr/></w:rPrChange>'
            '</w:rPr><w:t>C</w:t></w:r>'
            "</w:p>"
            '<w:p><w:pPr><w:rPr><w:del w:id="3" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>D</w:t></w:r></w:p>'
            '<w:p><w:pPr><w:rPr><w:ins w:id="4" w:author="Bob"/></w:rPr></w:pPr>'
            '<w:r><w:t>E</w:t></w:r></w:p>'
            "</w:body>"
        )

        revisions = list(document.revisions)

        assert [(revision.id, revision.author, revision.type) for revision in revisions] == [
            (5, "Carol", WD_REVISION_TYPE.PARAGRAPH_PROPERTY),
            (1, "Alice", WD_REVISION_TYPE.TEXT_INSERTION),
            (2, "Bob", WD_REVISION_TYPE.TEXT_DELETION),
            (6, "Carol", WD_REVISION_TYPE.CHARACTER_PROPERTY),
            (3, "Alice", WD_REVISION_TYPE.PARAGRAPH_MARK_DELETION),
            (4, "Bob", WD_REVISION_TYPE.PARAGRAPH_MARK_INSERTION),
        ]
        assert revisions[0].date is None

    def it_preserves_an_unselected_nested_insertion_when_accepting_outer_deletion(self):
        document = _document(
            "<w:body><w:p>"
            '<w:del w:id="1" w:author="Alice">'
            '<w:r><w:delText>old</w:delText></w:r>'
            '<w:ins w:id="2" w:author="Bob"><w:r><w:t>new</w:t></w:r></w:ins>'
            "</w:del></w:p></w:body>"
        )

        document.revisions.accept(1)

        assert document.paragraphs[0].text == "new"
        assert [(revision.id, revision.author) for revision in document.revisions] == [(2, "Bob")]

    def it_processes_a_batch_by_author_and_preserves_other_nested_and_adjacent_revisions(self):
        document = _document(
            "<w:body>"
            '<w:p>'
            '<w:ins w:id="1" w:author="Alice"><w:r><w:t>A</w:t></w:r></w:ins>'
            '<w:del w:id="2" w:author="Alice"><w:r><w:delText>a</w:delText></w:r></w:del>'
            '<w:ins w:id="3" w:author="Bob"><w:r><w:t>B</w:t></w:r></w:ins>'
            '<w:del w:id="4" w:author="Bob">'
            '<w:r><w:delText>x</w:delText></w:r>'
            '<w:ins w:id="5" w:author="Alice"><w:r><w:t>y</w:t></w:r></w:ins>'
            "</w:del>"
            "</w:p>"
            "</w:body>"
        )

        document.revisions.accept_all("Alice")

        assert document.paragraphs[0].text == "AB"
        assert [(revision.id, revision.author) for revision in document.revisions] == [
            (3, "Bob"),
            (4, "Bob"),
        ]

        document.revisions.reject_all("Bob")

        assert document.paragraphs[0].text == "Axy"
        assert len(document.revisions) == 0

    def it_accepts_a_deleted_paragraph_mark_with_surviving_properties(self):
        document = _document(
            "<w:body>"
            '<w:p><w:pPr><w:jc w:val="left"/>'
            '<w:rPr><w:del w:id="1" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>left</w:t></w:r></w:p>'
            '<w:p><w:pPr><w:jc w:val="right"/></w:pPr><w:r><w:t>right</w:t></w:r></w:p>'
            "</w:body>"
        )

        document.revisions.accept(1)

        assert len(document.paragraphs) == 1
        assert document.paragraphs[0].text == "leftright"
        assert document.paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.RIGHT

    def it_removes_left_properties_when_surviving_paragraph_has_no_properties(self):
        document = _document(
            "<w:body>"
            '<w:p><w:pPr><w:jc w:val="left"/>'
            '<w:rPr><w:del w:id="1" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>left</w:t></w:r></w:p>'
            '<w:p><w:r><w:t>right</w:t></w:r></w:p>'
            "</w:body>"
        )

        document.revisions.accept(1)

        assert len(document.paragraphs) == 1
        assert document.paragraphs[0].alignment is None

    def it_rejects_an_inserted_paragraph_mark_using_original_properties(self):
        document = _document(
            "<w:body>"
            '<w:p><w:pPr><w:jc w:val="left"/></w:pPr><w:r><w:t>left</w:t></w:r></w:p>'
            '<w:p><w:pPr><w:jc w:val="right"/>'
            '<w:rPr><w:ins w:id="1" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>right</w:t></w:r></w:p>'
            "</w:body>"
        )

        document.revisions.reject(1)

        assert len(document.paragraphs) == 1
        assert document.paragraphs[0].text == "leftright"
        assert document.paragraphs[0].alignment.name == "LEFT"

    def it_uses_original_properties_for_sequential_boundary_actions(self):
        document = _document(
            "<w:body>"
            '<w:p><w:pPr><w:jc w:val="left"/>'
            '<w:rPr><w:del w:id="1" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>left</w:t></w:r></w:p>'
            '<w:p><w:pPr><w:jc w:val="right"/>'
            '<w:rPr><w:ins w:id="2" w:author="Bob"/></w:rPr></w:pPr>'
            '<w:r><w:t>right</w:t></w:r></w:p>'
            "</w:body>"
        )

        document.revisions.accept(1)
        document.revisions.reject(2)

        assert len(document.paragraphs) == 1
        assert document.paragraphs[0].text == "leftright"
        assert document.paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.RIGHT

    def it_resolves_author_batch_boundary_conflict_using_pre_batch_identities(self):
        document = _document(
            "<w:body>"
            '<w:p><w:pPr><w:jc w:val="left"/>'
            '<w:rPr><w:del w:id="1" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>left</w:t></w:r></w:p>'
            '<w:p><w:pPr><w:jc w:val="right"/>'
            '<w:rPr><w:ins w:id="2" w:author="Bob"/></w:rPr></w:pPr>'
            '<w:r><w:t>right</w:t></w:r></w:p>'
            "</w:body>"
        )

        apply_revision_actions(document._element, {1: True, 2: False})

        assert len(document.paragraphs) == 1
        assert document.paragraphs[0].text == "leftright"
        assert document.paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.LEFT

    def it_restores_original_paragraph_and_run_properties(self):
        document = _document(
            "<w:body>"
            '<w:p>'
            '<w:pPr><w:jc w:val="right"/>'
            '<w:pPrChange w:id="1" w:author="Alice">'
            '<w:pPr><w:jc w:val="left"/></w:pPr></w:pPrChange>'
            "</w:pPr>"
            '<w:r><w:rPr><w:b/>'
            '<w:rPrChange w:id="2" w:author="Alice"><w:rPr><w:i/></w:rPr>'
            '</w:rPrChange></w:rPr>'
            '<w:t>text</w:t></w:r></w:p></w:body>'
        )

        document.revisions.reject((1, 2))

        paragraph = document.paragraphs[0]
        assert paragraph.alignment.name == "LEFT"
        assert paragraph.runs[0].bold is None
        assert paragraph.runs[0].italic is True

    def it_keeps_independent_paragraph_mark_font_revision_separate(self):
        document = _document(
            "<w:body>"
            '<w:p>'
            '<w:pPr><w:jc w:val="right"/><w:rPr><w:b/>'
            '<w:rPrChange w:id="2" w:author="Alice"><w:rPr><w:i/></w:rPr></w:rPrChange></w:rPr>'
            '<w:pPrChange w:id="1" w:author="Alice">'
            '<w:pPr><w:jc w:val="left"/></w:pPr></w:pPrChange>'
            "</w:pPr><w:r><w:t>text</w:t></w:r></w:p></w:body>"
        )

        document.revisions.reject((1, 2))

        assert document.paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.LEFT
        assert [(revision.id, revision.type) for revision in document.revisions] == []

    def it_round_trips_remaining_revisions_and_visible_text(self):
        document = _document(
            "<w:body>"
            '<w:p>'
            '<w:del w:id="1" w:author="Alice"><w:r><w:delText>old</w:delText></w:r></w:del>'
            '<w:ins w:id="2" w:author="Bob"><w:r><w:t>new</w:t></w:r></w:ins>'
            "</w:p></w:body>"
        )

        document.revisions.reject(1)
        reopened = _reopen(document)

        assert reopened.paragraphs[0].text == "oldnew"
        assert [(revision.id, revision.type) for revision in reopened.revisions] == [
            (2, WD_REVISION_TYPE.TEXT_INSERTION)
        ]

    def it_does_not_silently_accept_unsupported_revisions(self):
        document = _document(
            "<w:body>"
            '<w:p>'
            '<w:pPr><w:rPr><w:moveFrom w:id="1" w:author="Alice"/></w:rPr></w:pPr>'
            '<w:r><w:t>text</w:t></w:r></w:p></w:body>'
        )

        with pytest.raises(UnsupportedRevisionError):
            document.revisions.accept_all()

    def it_enumerates_a_revision_when_an_author_attribute_is_absent(self):
        document = _document(
            "<w:body><w:p>"
            '<w:ins w:id="1"><w:r><w:t>text</w:t></w:r></w:ins>'
            "</w:p></w:body>"
        )

        revision = next(iter(document.revisions))

        assert revision.id == 1
        assert revision.author is None
