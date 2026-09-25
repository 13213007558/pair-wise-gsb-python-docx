# pyright: reportPrivateUsage=false

"""Test suite for footnote and endnote support (`docx.notes` and friends)."""

from __future__ import annotations

import io

import pytest

from docx import Document
from docx.notes import Endnote, Footnote, Footnotes
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

IMAGE_PATH = "tests/test_files/python-icon.png"


def roundtrip(document) -> "Document":
    """Return a new Document loaded from the saved bytes of `document`."""
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return Document(buffer)


class DescribeFootnotes:
    """Behaviors of the `Footnotes` collection."""

    def it_starts_empty_apart_from_special_notes(self):
        footnotes = Document().footnotes

        assert isinstance(footnotes, Footnotes)
        assert len(footnotes) == 0
        assert list(footnotes) == []

    def it_can_add_a_footnote(self):
        footnotes = Document().footnotes

        footnote = footnotes.add_footnote("a footnote")

        assert isinstance(footnote, Footnote)
        assert footnote.note_id == 1
        assert footnote.text == "a footnote"
        assert len(footnotes) == 1

    def it_assigns_ids_that_do_not_collide_with_existing_notes(self):
        document = Document()
        first = document.add_paragraph("x").runs[0].add_footnote("one")

        document2 = roundtrip(document)
        second = document2.footnotes.add_footnote("two")

        assert first.note_id == 1
        assert second.note_id == 2

    def it_can_get_a_footnote_by_id(self):
        document = Document()
        footnote = document.footnotes.add_footnote("target")

        found = document.footnotes.get(footnote.note_id)
        assert found is not None
        assert found._note_elm is footnote._note_elm
        assert document.footnotes.get(999) is None

    def it_preserves_special_separator_notes(self):
        document = Document()
        document.footnotes.add_footnote()

        document2 = roundtrip(document)
        footnote_elms = document2.part.part_related_by(RT.FOOTNOTES).element

        types = [e.get(qn("w:type")) for e in footnote_elms]
        assert types == ["separator", "continuationSeparator", None]
        # -- special notes do not appear in the collection --
        assert len(document2.footnotes) == 1

    def it_refuses_to_remove_a_note_that_is_still_referenced(self):
        document = Document()
        footnote = document.add_paragraph("x").runs[0].add_footnote("refd")

        with pytest.raises(ValueError, match="still referenced"):
            document.footnotes.remove(footnote)

    def it_removes_a_note_that_is_no_longer_referenced(self):
        document = Document()
        footnote = document.add_paragraph("x").runs[0].add_footnote("doomed")
        footnote.references[0].remove()

        document.footnotes.remove(footnote)

        assert len(document.footnotes) == 0


class DescribeEndnotes:
    """Behaviors of the `Endnotes` collection."""

    def it_can_add_an_endnote(self):
        endnotes = Document().endnotes

        endnote = endnotes.add_endnote("an endnote")

        assert isinstance(endnote, Endnote)
        assert endnote.note_id == 1
        assert endnote.text == "an endnote"
        assert len(endnotes) == 1

    def it_preserves_special_separator_notes(self):
        document = Document()
        document.endnotes.add_endnote()

        document2 = roundtrip(document)
        endnote_elms = document2.part.part_related_by(RT.ENDNOTES).element

        types = [e.get(qn("w:type")) for e in endnote_elms]
        assert types == ["separator", "continuationSeparator", None]
        assert len(document2.endnotes) == 1


class DescribeRunNoteMethods:
    """Behaviors of the note-related `Run` methods."""

    def it_inserts_a_footnote_reference_at_the_run_boundary(self):
        document = Document()
        run = document.add_paragraph("anchor").runs[0]

        footnote = run.add_footnote("note text")

        reference_elms = document.element.body.xpath(
            f".//w:footnoteReference[@w:id='{footnote.note_id}']"
        )
        assert len(reference_elms) == 1
        # -- reference run directly follows the anchor run --
        assert reference_elms[0].getparent().getprevious() is run._r
        assert footnote.references[0].note_id == footnote.note_id

    def it_can_reference_the_same_footnote_more_than_once(self):
        document = Document()
        first = document.add_paragraph("one")
        second = document.add_paragraph().add_run("two")

        footnote = first.runs[0].add_footnote("shared")
        second.add_footnote_reference(footnote)

        assert len(footnote.references) == 2
        assert {r.note_id for r in footnote.references} == {footnote.note_id}

    def it_inserts_an_endnote_reference_at_the_run_boundary(self):
        document = Document()
        run = document.add_paragraph("anchor").runs[0]

        endnote = run.add_endnote("note text")

        reference_elms = document.element.body.xpath(
            f".//w:endnoteReference[@w:id='{endnote.note_id}']"
        )
        assert len(reference_elms) == 1
        assert reference_elms[0].getparent().getprevious() is run._r


class DescribeNote:
    """Behaviors of a single footnote or endnote."""

    def it_is_a_block_item_container(self):
        document = Document()
        footnote = document.footnotes.add_footnote("para one")

        footnote.add_paragraph("para two")
        table = footnote.add_table(2, 2, 3000000)
        table.cell(0, 0).text = "cell"

        assert len(footnote.paragraphs) == 2
        assert footnote.text == "para one\npara two"
        assert len(footnote.tables) == 1
        assert footnote.tables[0].cell(0, 0).text == "cell"

    def it_applies_the_default_note_paragraph_style(self):
        document = Document()
        footnote = document.footnotes.add_footnote()

        paragraph = footnote.add_paragraph()

        assert paragraph._p.style == "FootnoteText"

    def it_knows_its_note_type(self):
        document = Document()
        footnote = document.footnotes.add_footnote()

        assert footnote.note_type is None
        # -- special notes are still reachable by id --
        separator = document.footnotes.get(-1)
        assert separator is not None
        assert separator.note_type == "separator"

    def it_can_delete_itself_along_with_its_references(self):
        document = Document()
        run = document.add_paragraph("x").runs[0]
        footnote = run.add_footnote("gone")
        document.add_paragraph().add_run().add_footnote_reference(footnote)

        footnote.delete()

        assert len(document.footnotes) == 0
        assert document.element.body.xpath(".//w:footnoteReference") == []


class DescribeNoteReference:
    """Behaviors of a single note reference."""

    def it_removes_itself_and_its_otherwise_empty_run(self):
        document = Document()
        run = document.add_paragraph("x").runs[0]
        footnote = run.add_footnote("orphan")

        footnote.references[0].remove()

        assert footnote.references == []
        assert document.element.body.xpath(".//w:footnoteReference") == []
        # -- the anchor run is left in place --
        assert document.paragraphs[0].text == "x"

    def it_removes_a_specified_reference_only(self):
        document = Document()
        footnote = document.add_paragraph("a").runs[0].add_footnote("shared")
        document.add_paragraph().add_run("b").add_footnote_reference(footnote)

        footnote.references[0].remove()

        assert len(footnote.references) == 1


class DescribeNoteRelationships:
    """Relationships for note content belong to the notes part."""

    def it_relates_note_images_from_the_footnotes_part(self):
        document = Document()
        footnote = document.footnotes.add_footnote("with image")

        footnote.paragraphs[0].add_run().add_picture(IMAGE_PATH)

        footnotes_part = document.part.part_related_by(RT.FOOTNOTES)
        image_rels = [r for r in footnotes_part.rels.values() if r.reltype == RT.IMAGE]
        assert len(image_rels) == 1
        # -- no image relationship is added to the document part --
        assert [r for r in document.part.rels.values() if r.reltype == RT.IMAGE] == []

    def it_relates_note_hyperlinks_from_the_footnotes_part(self):
        document = Document()
        footnote = document.footnotes.add_footnote("with link")

        footnote.paragraphs[0].add_hyperlink("python", "https://python.org")

        footnotes_part = document.part.part_related_by(RT.FOOTNOTES)
        link_rels = [r for r in footnotes_part.rels.values() if r.reltype == RT.HYPERLINK]
        assert len(link_rels) == 1
        assert link_rels[0].target_ref == "https://python.org"
        assert link_rels[0].is_external
        # -- the hyperlink references a relationship of the footnotes part --
        hyperlink_elm = footnote.paragraphs[0].hyperlinks[0]._hyperlink
        assert hyperlink_elm.rId == link_rels[0].rId
        # -- no hyperlink relationship is added to the document part --
        assert [r for r in document.part.rels.values() if r.reltype == RT.HYPERLINK] == []


class DescribeNotesRoundTrip:
    """Round-trip sample: duplicate references, footnote image, endnote table."""

    def build_sample_document(self):
        document = Document()
        first = document.add_paragraph("A claim worth noting.")
        second = document.add_paragraph("The same note again.")
        third = document.add_paragraph("An endnote anchor.")

        # -- a footnote referenced from two positions, containing an image --
        footnote = first.runs[0].add_footnote("Footnote with picture.")
        second.runs[0].add_footnote_reference(footnote)
        footnote.paragraphs[0].add_run().add_picture(IMAGE_PATH)
        footnote.paragraphs[0].add_hyperlink("source", "https://example.com/src")

        # -- an endnote containing a table --
        endnote = third.runs[0].add_endnote("Endnote with table.")
        table = endnote.add_table(2, 2, 3000000)
        table.cell(0, 0).text = "r1c1"
        table.cell(1, 1).text = "r2c2"

        return document, footnote, endnote

    def it_round_trips_a_document_with_rich_notes(self):
        document, footnote, endnote = self.build_sample_document()
        footnote_id, endnote_id = footnote.note_id, endnote.note_id

        document2 = roundtrip(document)

        # -- notes and their ids survive the round-trip --
        assert len(document2.footnotes) == 1
        assert len(document2.endnotes) == 1
        footnote2 = document2.footnotes.get(footnote_id)
        endnote2 = document2.endnotes.get(endnote_id)
        assert footnote2 is not None
        assert endnote2 is not None

        # -- duplicate references survive --
        assert len(footnote2.references) == 2

        # -- footnote content, including its image relationship, survives --
        assert "Footnote with picture." in footnote2.text
        footnotes_part = document2.part.part_related_by(RT.FOOTNOTES)
        assert [r for r in footnotes_part.rels.values() if r.reltype == RT.IMAGE]
        assert [r for r in footnotes_part.rels.values() if r.reltype == RT.HYPERLINK]
        assert footnote2.paragraphs[0].hyperlinks[0].address == "https://example.com/src"

        # -- endnote table survives --
        assert len(endnote2.tables) == 1
        assert endnote2.tables[0].cell(0, 0).text == "r1c1"
        assert endnote2.tables[0].cell(1, 1).text == "r2c2"

        # -- special separator notes survive --
        types = [e.get(qn("w:type")) for e in footnotes_part.element]
        assert types == ["separator", "continuationSeparator", None]

    def it_preserves_unmodified_notes_and_relationships_on_resave(self):
        document, _, _ = self.build_sample_document()
        document2 = roundtrip(document)

        # -- add a new note after load; existing notes are not touched --
        new_footnote = document2.paragraphs[0].runs[0].add_footnote("added later")
        assert new_footnote.note_id == 2

        document3 = roundtrip(document2)

        assert len(document3.footnotes) == 2
        assert len(document3.endnotes) == 1
        footnotes_part = document3.part.part_related_by(RT.FOOTNOTES)
        # -- image and hyperlink relationships of the untouched note survive --
        assert len([r for r in footnotes_part.rels.values() if r.reltype == RT.IMAGE]) == 1
        assert len([r for r in footnotes_part.rels.values() if r.reltype == RT.HYPERLINK]) == 1
        original = document3.footnotes.get(1)
        assert original is not None
        assert "Footnote with picture." in original.text
        assert len(original.references) == 2
        endnotes_part = document3.part.part_related_by(RT.ENDNOTES)
        types = [e.get(qn("w:type")) for e in endnotes_part.element]
        assert types == ["separator", "continuationSeparator", None]
