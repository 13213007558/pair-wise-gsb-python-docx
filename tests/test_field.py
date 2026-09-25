"""Unit test suite for the docx.field module."""

from __future__ import annotations

import pytest

from docx import Document
from docx.exceptions import UnclosedFieldError
from docx.field import FieldPosition, iter_fields
from docx.oxml.ns import qn

from .unitutil.cxml import element


def body_with(*p_cxmls):
    """Return a w:body element containing paragraphs built from p_cxmls."""
    return element("w:body/(%s)" % ",".join(p_cxmls))


# -- a complex IF field whose instruction contains a nested REF field, with --
# -- begin/separate/end spread over many runs --
NESTED_IF_P = (
    "w:p/("
    "w:r/w:fldChar{w:fldCharType=begin},"
    'w:r/w:instrText" IF ",'
    "w:r/w:fldChar{w:fldCharType=begin},"
    'w:r/w:instrText" REF bk1 ",'
    "w:r/w:fldChar{w:fldCharType=separate},"
    'w:r/w:t"3",'
    "w:r/w:fldChar{w:fldCharType=end},"
    'w:r/w:instrText" = 3 yes no ",'
    "w:r/w:fldChar{w:fldCharType=separate},"
    'w:r/w:t"yes",'
    "w:r/w:fldChar{w:fldCharType=end}"
    ")"
)


class DescribeIterFields:
    """Unit-test suite for docx.field.iter_fields."""

    def it_recognizes_a_simple_field(self):
        body = body_with('w:p/w:fldSimple{w:instr= REF bk1 }/w:r/w:t"42"')

        fields = iter_fields(body)

        assert len(fields) == 1
        fld = fields[0]
        assert fld.is_simple is True
        assert fld.is_complex is False
        assert fld.instruction == "REF bk1"
        assert fld.result == "42"
        assert fld.is_broken is False
        assert fld.parent is None

    def it_recognizes_a_complex_field_spread_across_runs(self):
        body = body_with(
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" DOCVARIABLE ",'
            'w:r/w:instrText" name ",'
            "w:r/w:fldChar{w:fldCharType=separate},"
            'w:r/w:t"Alice",'
            "w:r/w:fldChar{w:fldCharType=end}"
            ")"
        )

        (fld,) = iter_fields(body)

        assert fld.is_complex is True
        assert fld.instruction == "DOCVARIABLE  name"
        assert fld.result == "Alice"

    def it_parses_nested_fields_across_runs(self):
        body = body_with(NESTED_IF_P)

        (outer,) = iter_fields(body)

        assert outer.instruction == "IF  = 3 yes no"
        assert outer.result == "yes"
        assert len(outer.children) == 1
        inner = outer.children[0]
        assert inner.parent is outer
        assert inner.instruction == "REF bk1"
        assert inner.result == "3"

    def it_includes_nested_field_results_in_the_outer_result_text(self):
        body = body_with(
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" QUOTE ",'
            "w:r/w:fldChar{w:fldCharType=separate},"
            'w:r/w:t"a=",'
            'w:fldSimple{w:instr= REF n }/w:r/w:t"7",'
            "w:r/w:fldChar{w:fldCharType=end}"
            ")"
        )

        (outer,) = iter_fields(body)

        assert outer.result == "a=7"
        assert [c.instruction for c in outer.children] == ["REF n"]

    def it_reads_a_result_containing_breaks_and_a_hyperlink(self):
        body = body_with(
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" DOCVARIABLE x ",'
            "w:r/w:fldChar{w:fldCharType=separate},"
            'w:r/(w:t"line1",w:br,w:t"line2"),'
            'w:hyperlink{w:anchor=bm}/w:r/w:t"click",'
            "w:r/w:fldChar{w:fldCharType=end}"
            ")"
        )

        (fld,) = iter_fields(body)

        assert fld.result == "line1\nline2click"

    def it_reports_unclosed_fields_with_position_without_swallowing_content(self):
        body = body_with(
            'w:p/w:r/w:t"intro",'
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" REF z "'
            "),"
            'w:p/w:r/w:t"tail"'
        )

        (fld,) = iter_fields(body)

        assert fld.is_broken is True
        assert fld.position == FieldPosition(paragraph_index=1, run_index=0)
        assert fld.result is None
        # -- the rest of the story is not swallowed as the field result --
        texts = [t.text for t in body.iter(qn("w:t"))]
        assert texts == ["intro", "tail"]

    def it_parses_each_story_independently(self):
        body = element(
            "w:body/("
            'w:p/w:fldSimple{w:instr= REF A }/w:r/w:t"1",'
            'w:tbl/w:tr/w:tc/w:p/w:fldSimple{w:instr= REF B }/w:r/w:t"2"'
            ")"
        )

        body_fields = iter_fields(body)
        tc = body.find(".//" + qn("w:tc"))
        cell_fields = iter_fields(tc)

        assert [f.instruction for f in body_fields] == ["REF A"]
        assert [f.instruction for f in cell_fields] == ["REF B"]

    def it_does_not_let_an_unclosed_field_swallow_a_table_cell_story(self):
        body = element(
            "w:body/("
            "w:p/w:r/w:fldChar{w:fldCharType=begin},"
            'w:tbl/w:tr/w:tc/w:p/w:r/w:t"cell text"'
            ")"
        )

        (fld,) = iter_fields(body)

        assert fld.is_broken is True
        assert fld.result is None
        tc = body.find(".//" + qn("w:tc"))
        cell_texts = [t.text for t in tc.iter(qn("w:t"))]
        assert cell_texts == ["cell text"]


class DescribeReplaceResult:
    """Unit-test suite for `Field.replace_result`."""

    def it_replaces_the_result_of_a_complex_field(self):
        body = body_with(NESTED_IF_P)
        (outer,) = iter_fields(body)

        outer.replace_result("no")

        assert outer.result == "no"
        # -- instruction and nested field in the instruction region survive --
        assert outer.instruction == "IF  = 3 yes no"
        assert len(outer.children) == 1
        inner = outer.children[0]
        assert inner.instruction == "REF bk1"
        assert inner.result == "3"
        # -- and the XML still holds the nested field-char runs --
        fld_char_types = [fc.get(qn("w:fldCharType")) for fc in body.iter(qn("w:fldChar"))]
        assert fld_char_types == ["begin", "begin", "separate", "end", "separate", "end"]

    def it_preserves_bookmarks_and_runs_outside_the_field(self):
        body = body_with(
            "w:p/("
            "w:bookmarkStart{w:id=0,w:name=bk1},"
            'w:r/(w:rPr/w:b,w:t"before"),'
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" REF x ",'
            "w:r/w:fldChar{w:fldCharType=separate},"
            'w:r/w:t"old",'
            "w:r/w:fldChar{w:fldCharType=end},"
            'w:r/(w:rPr/w:i,w:t"after"),'
            "w:bookmarkEnd{w:id=0}"
            ")"
        )
        (fld,) = iter_fields(body)

        fld.replace_result("new")

        p = body.find(qn("w:p"))
        assert p.find(qn("w:bookmarkStart")) is not None
        assert p.find(qn("w:bookmarkEnd")) is not None
        texts = [t.text for t in p.iter(qn("w:t"))]
        assert texts == ["before", "new", "after"]
        # -- formatting of the surrounding runs is untouched --
        runs = p.findall(qn("w:r"))
        assert runs[0].find(qn("w:rPr")) is not None
        assert runs[-1].find(qn("w:rPr")) is not None

    def it_carries_over_the_formatting_of_the_first_result_run(self):
        body = body_with(
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" REF x ",'
            "w:r/w:fldChar{w:fldCharType=separate},"
            'w:r/(w:rPr/w:b,w:t"old"),'
            "w:r/w:fldChar{w:fldCharType=end}"
            ")"
        )
        (fld,) = iter_fields(body)

        fld.replace_result("new")

        assert fld.result == "new"
        # -- the new run sits just before the end field-char run and carries --
        # -- the bold formatting of the replaced result run --
        new_r = body.find(qn("w:p")).findall(qn("w:r"))[-2]
        r_pr = new_r.find(qn("w:rPr"))
        assert r_pr is not None
        assert r_pr.find(qn("w:b")) is not None

    def it_replaces_the_result_of_a_simple_field(self):
        body = body_with('w:p/w:fldSimple{w:instr= REF y }/w:r/w:t"42"')
        (fld,) = iter_fields(body)

        fld.replace_result("43")

        fld_simple = body.find(qn("w:p")).find(qn("w:fldSimple"))
        assert fld_simple.get(qn("w:instr")) == "REF y "
        assert fld.result == "43"

    def it_inserts_a_separator_when_the_field_has_no_cached_result(self):
        body = body_with(
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" REF q ",'
            "w:r/w:fldChar{w:fldCharType=end}"
            ")"
        )
        (fld,) = iter_fields(body)
        assert fld.result is None

        fld.replace_result("v")

        assert fld.result == "v"
        assert fld.instruction == "REF q"
        fld_char_types = [fc.get(qn("w:fldCharType")) for fc in body.iter(qn("w:fldChar"))]
        assert fld_char_types == ["begin", "separate", "end"]

    def it_writes_newlines_as_break_elements(self):
        body = body_with(
            "w:p/("
            "w:r/w:fldChar{w:fldCharType=begin},"
            'w:r/w:instrText" REF q ",'
            "w:r/w:fldChar{w:fldCharType=separate},"
            'w:r/w:t"old",'
            "w:r/w:fldChar{w:fldCharType=end}"
            ")"
        )
        (fld,) = iter_fields(body)

        fld.replace_result("x\ny")

        assert fld.result == "x\ny"
        assert body.find(qn("w:p")).find(qn("w:r")).find(qn("w:br")) is None
        brs = list(body.iter(qn("w:br")))
        assert len(brs) == 1

    def it_raises_on_an_unclosed_field(self):
        body = body_with(
            "w:p/(" "w:r/w:fldChar{w:fldCharType=begin}," 'w:r/w:instrText" REF z "' ")"
        )
        (fld,) = iter_fields(body)

        with pytest.raises(UnclosedFieldError):
            fld.replace_result("x")


class DescribeStoryFields:
    """Integration tests for `.fields` on document, cell, and header stories."""

    def it_parses_body_cell_and_header_stories_independently(self):
        document = Document()
        document.element.body.insert(0, element(NESTED_IF_P))
        table = document.add_table(1, 1)
        cell = table.cell(0, 0)
        cell._tc.append(element('w:p/w:fldSimple{w:instr= REF C }/w:r/w:t"2"'))
        header = document.sections[0].header
        header.is_linked_to_previous = False
        header._element.append(element('w:p/w:fldSimple{w:instr= PAGE }/w:r/w:t"3"'))

        body_fields = document.fields
        cell_fields = cell.fields
        header_fields = header.fields

        assert [f.instruction for f in body_fields] == ["IF  = 3 yes no"]
        assert body_fields[0].children[0].instruction == "REF bk1"
        assert [f.instruction for f in cell_fields] == ["REF C"]
        assert [f.instruction for f in header_fields] == ["PAGE"]

    def it_replaces_a_result_through_the_document_view(self):
        document = Document()
        document.element.body.insert(0, element(NESTED_IF_P))

        document.fields[0].replace_result("no")

        assert document.fields[0].result == "no"
