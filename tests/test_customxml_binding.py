"""Tests for content controls bound to custom XML data-stores."""

from __future__ import annotations

import datetime as dt
import io

import pytest

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml
from docx.sdt import (
    INCOMPATIBLE_CONTROL,
    STORE_NOT_FOUND,
    UNSUPPORTED_CONTROL,
    XPATH_INVALID,
    XPATH_MULTIPLE_MATCHES,
    XPATH_NO_MATCH,
    BindingConflictError,
)

NS_A = "http://example.com/store-a"
NS_B = "http://example.com/store-b"

STORE_A_XML = (
    '<root xmlns="%s">'
    "<name>Old Name</name>"
    "<date>2026-01-15</date>"
    "<flag>false</flag>"
    "</root>" % NS_A
)
STORE_B_XML = '<root xmlns="%s"><name>Other Name</name></root>' % NS_B

NSDECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:x="http://example.com/extension"'
)

NAME_XPATH = "/ns0:root/ns0:name"
DATE_XPATH = "/ns0:root/ns0:date"
FLAG_XPATH = "/ns0:root/ns0:flag"

PREFIX_MAP_A = "xmlns:ns0='%s'" % NS_A
PREFIX_MAP_B = "xmlns:ns0='%s'" % NS_B


def _text_sdt(store_id, sdt_id, text, prefix_map=PREFIX_MAP_A, xpath=NAME_XPATH):
    return (
        "<w:sdt %s>"
        "<w:sdtPr>"
        '<w:alias w:val="Name Control"/>'
        '<w:tag w:val="name-tag"/>'
        '<w:id w:val="%d"/>'
        '<w:lock w:val="sdtContentLocked"/>'
        '<w:dataBinding w:storeItemID="%s" w:xpath="%s" w:prefixMappings="%s"/>'
        '<x:extension x:track="keep-me"/>'
        "<w:text/>"
        "</w:sdtPr>"
        "<w:sdtContent>"
        "<w:r><w:rPr><w:b/></w:rPr><w:t>%s</w:t></w:r>"
        "</w:sdtContent>"
        "</w:sdt>" % (NSDECL, sdt_id, store_id, xpath, prefix_map, text)
    )


def _date_sdt(store_id, sdt_id, text, date_format, prefix_map=PREFIX_MAP_A):
    return (
        "<w:sdt %s>"
        "<w:sdtPr>"
        '<w:id w:val="%d"/>'
        '<w:dataBinding w:storeItemID="%s" w:xpath="%s" w:prefixMappings="%s"/>'
        "<w:date>"
        '<w:dateFormat w:val="%s"/>'
        '<w:calendar w:val="gregorian"/>'
        '<w:lid w:val="zh-CN"/>'
        "</w:date>"
        "</w:sdtPr>"
        "<w:sdtContent>"
        "<w:r><w:t>%s</w:t></w:r>"
        "</w:sdtContent>"
        "</w:sdt>" % (NSDECL, sdt_id, store_id, DATE_XPATH, prefix_map, date_format, text)
    )


def _checkbox_sdt(store_id, sdt_id, checked, prefix_map=PREFIX_MAP_A):
    glyph, val = ("\u2612", "1") if checked else ("\u2610", "0")
    return (
        "<w:sdt %s>"
        "<w:sdtPr>"
        '<w:id w:val="%d"/>'
        '<w:dataBinding w:storeItemID="%s" w:xpath="%s" w:prefixMappings="%s"/>'
        "<w14:checkbox>"
        '<w14:checked w14:val="%s"/>'
        '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/>'
        "</w14:checkbox>"
        "</w:sdtPr>"
        "<w:sdtContent>"
        '<w:r><w:rPr><w:rFonts w:ascii="MS Gothic" w:hAnsi="MS Gothic"/></w:rPr>'
        "<w:t>%s</w:t></w:r>"
        "</w:sdtContent>"
        "</w:sdt>" % (NSDECL, sdt_id, store_id, FLAG_XPATH, prefix_map, val, glyph)
    )


def _append_sdt(container_p, sdt_xml):
    container_p._p.append(parse_xml(sdt_xml))


@pytest.fixture
def bound_doc():
    """A document with two stores and controls in body, header, and a table.

    Store A holds name/date/flag items. The name item is bound to a body control
    and a header control; the date item is bound to a body control and a
    table-cell control; the flag item is bound to a body checkbox. Store B holds
    a same-named item bound to a separate body control.
    """
    doc = Document()
    store_a = doc.custom_xml_stores.add(STORE_A_XML)
    store_b = doc.custom_xml_stores.add(STORE_B_XML)
    id_a, id_b = store_a.store_item_id, store_b.store_item_id

    _append_sdt(doc.add_paragraph(), _text_sdt(id_a, 101, "Old Name"))
    _append_sdt(doc.add_paragraph(), _date_sdt(id_a, 102, "2026-01-15", "yyyy-MM-dd"))
    _append_sdt(doc.add_paragraph(), _checkbox_sdt(id_a, 103, checked=False))
    _append_sdt(doc.add_paragraph(), _text_sdt(id_b, 104, "Other Name", PREFIX_MAP_B))

    header = doc.sections[0].header
    header.is_linked_to_previous = False
    _append_sdt(header.paragraphs[0], _text_sdt(id_a, 201, "Old Name"))

    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    _append_sdt(
        cell.paragraphs[0],
        _date_sdt(id_a, 301, "2026\u5e741\u670815\u65e5", "yyyy\u5e74MM\u6708dd\u65e5"),
    )

    return doc


def _store_text(store, xpath, ns):
    return store.find(xpath, ns)[0].text


class DescribeFindBoundControls:
    def it_finds_controls_across_body_header_and_tables(self, bound_doc):
        controls = bound_doc.find_bound_controls()
        assert sorted(c.control_id for c in controls) == ["101", "102", "103", "104", "201", "301"]
        part_names = {c.part_name for c in controls}
        assert "/word/document.xml" in part_names
        assert "/word/header1.xml" in part_names

    def it_filters_by_store_item_id(self, bound_doc):
        store_a, store_b = bound_doc.custom_xml_stores
        id_a, id_b = store_a.store_item_id, store_b.store_item_id
        assert id_a != id_b
        ids_a = {c.control_id for c in bound_doc.find_bound_controls(store_item_id=id_a)}
        assert ids_a == {"101", "102", "103", "201", "301"}
        ids_b = {c.control_id for c in bound_doc.find_bound_controls(store_item_id=id_b)}
        assert ids_b == {"104"}

    def it_filters_by_store_item_id_case_insensitively(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        lowered = store_a.store_item_id.lower()
        assert len(bound_doc.find_bound_controls(store_item_id=lowered)) == 5

    def it_filters_by_xpath(self, bound_doc):
        controls = bound_doc.find_bound_controls(xpath=DATE_XPATH)
        assert {c.control_id for c in controls} == {"102", "301"}
        assert all(c.control_type == "date" for c in controls)


class DescribeUpdateBoundControls:
    def it_syncs_text_across_body_and_header(self, bound_doc):
        store_a, store_b = bound_doc.custom_xml_stores
        id_a = store_a.store_item_id

        result = bound_doc.update_bound_controls(id_a, NAME_XPATH, "New Name")

        assert _store_text(store_a, NAME_XPATH, {"ns0": NS_A}) == "New Name"
        by_id = {c.control_id: c for c in bound_doc.find_bound_controls()}
        assert by_id["101"].displayed_text == "New Name"
        assert by_id["201"].displayed_text == "New Name"
        assert sorted(c.control_id for c in result.updated_controls) == ["101", "201"]
        assert "/word/header1.xml" in " ".join(result.locations)
        # -- same-xpath control bound to the other store is untouched --
        assert _store_text(store_b, NAME_XPATH, {"ns0": NS_B}) == "Other Name"
        assert by_id["104"].displayed_text == "Other Name"

    def it_formats_date_display_per_control(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        id_a = store_a.store_item_id

        bound_doc.update_bound_controls(id_a, DATE_XPATH, dt.date(2026, 9, 25))

        assert _store_text(store_a, DATE_XPATH, {"ns0": NS_A}) == "2026-09-25"
        by_id = {c.control_id: c for c in bound_doc.find_bound_controls()}
        assert by_id["102"].displayed_text == "2026-09-25"
        assert by_id["301"].displayed_text == "2026\u5e7409\u670825\u65e5"

    def it_updates_checkbox_state_and_display(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        id_a = store_a.store_item_id

        bound_doc.update_bound_controls(id_a, FLAG_XPATH, True)

        assert _store_text(store_a, FLAG_XPATH, {"ns0": NS_A}) == "true"
        (control,) = [c for c in bound_doc.find_bound_controls() if c.control_id == "103"]
        assert control.displayed_text == "\u2612"
        checkbox = control._sdt.find(qn("w:sdtPr")).find(qn("w14:checkbox"))
        assert checkbox.find(qn("w14:checked")).get(qn("w14:val")) == "1"

    def it_preserves_styles_locks_and_unknown_extensions(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        id_a = store_a.store_item_id

        bound_doc.update_bound_controls(id_a, NAME_XPATH, "New Name")

        (control,) = [c for c in bound_doc.find_bound_controls() if c.control_id == "101"]
        sdtPr = control._sdt.find(qn("w:sdtPr"))
        assert sdtPr.find(qn("w:lock")).get(qn("w:val")) == "sdtContentLocked"
        assert sdtPr.find(qn("w:alias")).get(qn("w:val")) == "Name Control"
        extension = sdtPr.find("{http://example.com/extension}extension")
        assert extension is not None
        assert extension.get("{http://example.com/extension}track") == "keep-me"
        run = control._sdt.find(qn("w:sdtContent")).find(qn("w:r"))
        assert run.find(qn("w:rPr")).find(qn("w:b")) is not None
        assert run.find(qn("w:t")).text == "New Name"

    def it_updates_store_when_no_controls_are_bound(self, bound_doc):
        store_c = bound_doc.custom_xml_stores.add(
            '<root xmlns="%s"><unbound>x</unbound></root>' % NS_A
        )

        result = bound_doc.update_bound_controls(
            store_c.store_item_id, "/ns0:root/ns0:unbound", "y", {"ns0": NS_A}
        )

        assert _store_text(store_c, "/ns0:root/ns0:unbound", {"ns0": NS_A}) == "y"
        assert result.updated_controls == []


class DescribeSaveAndReopen:
    def it_keeps_bindings_and_store_isolation_after_roundtrip(self, bound_doc):
        store_a, store_b = bound_doc.custom_xml_stores
        id_a, id_b = store_a.store_item_id, store_b.store_item_id
        bound_doc.update_bound_controls(id_a, NAME_XPATH, "New Name")
        bound_doc.update_bound_controls(id_a, DATE_XPATH, dt.date(2026, 9, 25))
        bound_doc.update_bound_controls(id_a, FLAG_XPATH, True)

        stream = io.BytesIO()
        bound_doc.save(stream)
        stream.seek(0)
        doc = Document(stream)

        stores = {s.store_item_id: s for s in doc.custom_xml_stores}
        assert set(stores) == {id_a, id_b}
        assert _store_text(stores[id_a], NAME_XPATH, {"ns0": NS_A}) == "New Name"
        assert _store_text(stores[id_a], DATE_XPATH, {"ns0": NS_A}) == "2026-09-25"
        assert _store_text(stores[id_a], FLAG_XPATH, {"ns0": NS_A}) == "true"
        assert _store_text(stores[id_b], NAME_XPATH, {"ns0": NS_B}) == "Other Name"

        by_id = {c.control_id: c for c in doc.find_bound_controls()}
        assert sorted(by_id) == ["101", "102", "103", "104", "201", "301"]
        assert by_id["101"].store_item_id == id_a
        assert by_id["201"].store_item_id == id_a
        assert by_id["104"].store_item_id == id_b
        assert by_id["101"].displayed_text == "New Name"
        assert by_id["201"].displayed_text == "New Name"
        assert by_id["301"].displayed_text == "2026\u5e7409\u670825\u65e5"
        assert by_id["103"].displayed_text == "\u2612"
        # -- the control bound to the other store did not follow the update --
        assert by_id["104"].displayed_text == "Other Name"


class DescribeBindingConflicts:
    def it_raises_store_not_found_for_an_unknown_store(self, bound_doc):
        with pytest.raises(BindingConflictError) as excinfo:
            bound_doc.update_bound_controls("{DEADBEEF-0000-0000-0000-000000000000}", NAME_XPATH, "x")
        assert [c.reason for c in excinfo.value.conflicts] == [STORE_NOT_FOUND]

    def it_raises_and_writes_nothing_when_xpath_matches_nothing(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        id_a = store_a.store_item_id

        with pytest.raises(BindingConflictError) as excinfo:
            bound_doc.update_bound_controls(id_a, "/ns0:root/ns0:missing", "x", {"ns0": NS_A})

        assert [c.reason for c in excinfo.value.conflicts] == [XPATH_NO_MATCH]
        assert _store_text(store_a, NAME_XPATH, {"ns0": NS_A}) == "Old Name"

    def it_raises_when_xpath_cannot_be_evaluated(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        id_a = store_a.store_item_id

        with pytest.raises(BindingConflictError) as excinfo:
            bound_doc.update_bound_controls(id_a, "/ns0:root/ns0:missing", "x")

        assert [c.reason for c in excinfo.value.conflicts] == [XPATH_INVALID]
        assert _store_text(store_a, NAME_XPATH, {"ns0": NS_A}) == "Old Name"

    def it_raises_and_writes_nothing_when_xpath_matches_many(self, bound_doc):
        store_c = bound_doc.custom_xml_stores.add(
            '<root xmlns="%s"><name>a</name><name>b</name></root>' % NS_A
        )

        with pytest.raises(BindingConflictError) as excinfo:
            bound_doc.update_bound_controls(
                store_c.store_item_id, NAME_XPATH, "x", {"ns0": NS_A}
            )

        assert [c.reason for c in excinfo.value.conflicts] == [XPATH_MULTIPLE_MATCHES]
        nodes = store_c.find(NAME_XPATH, {"ns0": NS_A})
        assert [n.text for n in nodes] == ["a", "b"]

    def it_raises_locatable_conflict_for_incompatible_controls(self, bound_doc):
        store_d = bound_doc.custom_xml_stores.add(
            '<root xmlns="%s"><flag>false</flag></root>' % NS_A
        )
        id_d = store_d.store_item_id
        _append_sdt(bound_doc.add_paragraph(), _checkbox_sdt(id_d, 401, checked=False))
        _append_sdt(
            bound_doc.add_paragraph(),
            _text_sdt(id_d, 402, "false", xpath=FLAG_XPATH),
        )

        with pytest.raises(BindingConflictError) as excinfo:
            bound_doc.update_bound_controls(id_d, FLAG_XPATH, True)

        conflicts = excinfo.value.conflicts
        assert [c.reason for c in conflicts] == [INCOMPATIBLE_CONTROL]
        assert conflicts[0].location == "/word/document.xml#sdt[402]"
        # -- atomic: neither the stored XML nor any display was modified --
        assert _store_text(store_d, FLAG_XPATH, {"ns0": NS_A}) == "false"
        by_id = {c.control_id: c for c in bound_doc.find_bound_controls()}
        assert by_id["401"].displayed_text == "\u2610"
        assert by_id["402"].displayed_text == "false"

    def it_raises_locatable_conflict_for_unsupported_controls(self, bound_doc):
        store_a = bound_doc.custom_xml_stores[0]
        id_a = store_a.store_item_id
        rich_text_sdt = _text_sdt(id_a, 501, "Old Name").replace("<w:text/>", "<w:richText/>")
        _append_sdt(bound_doc.add_paragraph(), rich_text_sdt)

        with pytest.raises(BindingConflictError) as excinfo:
            bound_doc.update_bound_controls(id_a, NAME_XPATH, "New Name")

        conflicts = excinfo.value.conflicts
        assert [c.reason for c in conflicts] == [UNSUPPORTED_CONTROL]
        assert conflicts[0].location == "/word/document.xml#sdt[501]"
        assert _store_text(store_a, NAME_XPATH, {"ns0": NS_A}) == "Old Name"

    def it_accepts_a_value_all_bound_control_types_can_display(self, bound_doc):
        store_d = bound_doc.custom_xml_stores.add(
            '<root xmlns="%s"><flag>false</flag></root>' % NS_A
        )
        id_d = store_d.store_item_id
        _append_sdt(bound_doc.add_paragraph(), _checkbox_sdt(id_d, 401, checked=False))
        _append_sdt(
            bound_doc.add_paragraph(),
            _text_sdt(id_d, 402, "false", xpath=FLAG_XPATH),
        )

        bound_doc.update_bound_controls(id_d, FLAG_XPATH, "true")

        assert _store_text(store_d, FLAG_XPATH, {"ns0": NS_A}) == "true"
        by_id = {c.control_id: c for c in bound_doc.find_bound_controls()}
        assert by_id["401"].displayed_text == "\u2612"
        assert by_id["402"].displayed_text == "true"
